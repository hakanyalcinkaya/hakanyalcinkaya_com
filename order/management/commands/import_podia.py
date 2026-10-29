"""Podia fatura dışa aktarımını (xlsx) siparişlere ve kurs kayıtlarına aktarır.

Örnek:
    python manage.py import_podia podia-invoices.xlsx \\
        --map "React 18.x=react" --map "Python ve Django=python-django" --map "İki Eğitmen=sifirdan-python" --dry-run
"""

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from order.podia import REQUIRED_COLUMNS, PodiaImportError, course_rules, import_rows


class Command(BaseCommand):
    help = "Podia satışlarını (fatura xlsx) siparişlere ve kurs kayıtlarına aktarır; tekrar çalıştırılabilir."

    def add_arguments(self, parser):
        parser.add_argument("path", help="Podia'dan indirilen fatura dosyası (.xlsx)")
        parser.add_argument(
            "--map",
            action="append",
            default=[],
            dest="rules",
            help="Ürün adında geçen metin=kurs-slug (her kurs için bir kez).",
        )
        parser.add_argument("--dry-run", action="store_true", help="Hiçbir şey kaydetmeden raporla.")

    def handle(self, path, rules, dry_run, **options):
        try:
            sheet = load_workbook(path, read_only=True, data_only=True).active
        except (OSError, ValueError) as error:
            raise CommandError(f"Dosya okunamadı: {error}") from error
        values = sheet.iter_rows(values_only=True)
        header = [str(cell or "").strip() for cell in next(values, [])]
        missing = [column for column in REQUIRED_COLUMNS if column not in header]
        if missing:
            raise CommandError(f"Dosyada eksik sütun: {', '.join(missing)}")
        rows = [dict(zip(header, row)) for row in values if any(row)]
        try:
            report = import_rows(rows, course_rules(rules), dry_run=dry_run)
        except PodiaImportError as error:
            raise CommandError(str(error)) from error

        title = "DENEME (hiçbir şey kaydedilmedi)" if dry_run else "AKTARIM TAMAMLANDI"
        self.stdout.write(self.style.WARNING(title) if dry_run else self.style.SUCCESS(title))
        self.stdout.write(f"Satır: {len(rows)}")
        self.stdout.write(f"Sipariş: {report.orders_created} yeni, {report.orders_skipped} daha önce aktarılmış")
        self.stdout.write(f"Kullanıcı: {report.users_created} yeni, {report.users_existing} var olan")
        self.stdout.write(f"Kurs kaydı: {report.enrollments_created} yeni, {report.enrollments_existing} zaten vardı")
        if report.skipped_status:
            self.stdout.write(f"Tamamlanmamış (atlandı): {report.skipped_status}")
        for item, count in report.unmatched_items.items():
            self.stdout.write(self.style.ERROR(f"Eşleşmeyen ürün ({count} satır): {item}"))
        for error in report.errors:
            self.stdout.write(self.style.ERROR(error))
